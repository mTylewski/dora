"""Step 7: draft demand letter (.docx) from the template.

The model writes only the justification paragraph, from anonymized facts
prepared by code. Party data is entered by the lawyer and inserted locally;
it is never sent to the API.
"""

from __future__ import annotations

import copy
import datetime as dt
import json
import re
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any

import yaml
from docx import Document as DocxDocument
from docx.table import _Row
from docx.text.paragraph import Paragraph
from pydantic import BaseModel, ConfigDict

from calculator.core import EVENT_LABELS, PenaltyResult, fmt_date, fmt_percent
from calculator.money import format_pln, round_grosze
from pipeline.calculate import CalculationOutcome
from pipeline.llm import LLMClient, call_structured
from pipeline.schema import Extraction, strict_schema
from pipeline.settings import CONFIG_DIR, TEMPLATES_DIR, StepSettings, load_step, read_prompt

TEMPLATE = TEMPLATES_DIR / "wezwanie_do_zaplaty.docx"
DRAFT_BANNER = "PROJEKT — do weryfikacji i podpisu przez pełnomocnika"
MISSING = "[UZUPEŁNIJ: {}]"
_PLACEHOLDER = re.compile(r"\{\{(\w+)\}\}")
_NUMBER = re.compile(r"\d[\d  .,]*\d|\d")


@dataclass(frozen=True)
class LetterParties:
    """Letter data entered or confirmed by the lawyer in the wizard; stays on the server.

    Empty firm fields fall back to config/letter.yaml (CLI); the web app fills them from the firm settings.
    """

    client_name: str = ""  # one or more clients, e.g. "Jan Nowak i Anna Nowak"
    client_address: str = ""
    developer_name: str = ""
    developer_address: str = ""
    agreement_reference: str = ""  # e.g. "umową deweloperską z dnia ... (Rep. A nr ...)"
    bank_account: str = ""
    case_number: str = ""
    lawyer_name: str = ""
    lawyer_title: str = ""
    lawyer_registry: str = ""
    agreement_date: str = ""  # DD.MM.RRRR
    agreement_deed: str = ""  # "Rep. A nr ..."
    firm_name: str = ""
    firm_address: str = ""
    city: str = ""
    payment_deadline_days: int | None = None
    legal_basis: str = ""
    clients: tuple[tuple[str, str], ...] = ()  # (name, address) per client, for the intro sentence


@dataclass
class LetterResult:
    path: Path
    total_grosze: int
    warnings: list[str] = field(default_factory=list)


class Justification(BaseModel):
    model_config = ConfigDict(extra="forbid")

    justification: str


def load_letter_config(config_dir: Path | None = None) -> dict[str, Any]:
    with ((config_dir or CONFIG_DIR) / "letter.yaml").open(encoding="utf-8") as fh:
        return yaml.safe_load(fh)


def demanded_grosze(p: PenaltyResult, capped: bool) -> int:
    value = p.capped_grosze if capped else p.uncapped_grosze
    return value or 0


def letter_facts(ex: Extraction, outcome: CalculationOutcome, capped: bool, legal_basis: str) -> dict[str, Any]:
    """Facts for the model: dates, numbers and the penalty clause; no party data."""
    assert outcome.result is not None
    penalties = []
    for p in outcome.result.penalties:
        if p.uncapped_grosze is None:
            continue
        penalties.append({
            "event": EVENT_LABELS[p.event],
            "deadline": fmt_date(p.effective_deadline.deadline),
            "deadline_source": p.effective_deadline.describe(),
            "actual_date": None if p.ongoing else fmt_date(p.end_date),
            "still_not_performed_as_of": fmt_date(p.end_date) if p.ongoing else None,
            "days_of_delay": p.days,
            "amount_demanded": format_pln(demanded_grosze(p, capped)),
            "cap_applied": capped and p.cap_reached,
        })
    rate = ex.penalty_rate.value
    return {
        "penalty_clause_quote": ex.penalty_rate.quote or ex.penalty_basis.quote,
        "penalty_rate": (fmt_percent(rate) + " ceny za każdy dzień")
        if rate is not None and ex.penalty_basis.value == "percent_of_price_per_day"
        else (format_pln(round_grosze(rate * 100)) + " za każdy dzień" if rate is not None else None),
        "price": format_pln(round_grosze(ex.price_gross.value * 100)) if ex.price_gross.value is not None else None,
        "penalties": penalties,
        "total_demanded": format_pln(sum(demanded_grosze(p, capped) for p in outcome.result.penalties)),
        "legal_basis": legal_basis,
    }


def unsupported_numbers(text: str, facts: dict[str, Any]) -> list[str]:
    """Numbers in the model's text that do not appear in the facts."""
    def norm(s: str) -> str:
        return re.sub(r"[  ]", "", s).rstrip(".,")

    known = {norm(m) for m in _NUMBER.findall(json.dumps(facts, ensure_ascii=False))}
    return sorted({m for m in (norm(x) for x in _NUMBER.findall(text)) if m and m not in known})


def draft_justification(
    client: LLMClient, case_id: str, facts: dict[str, Any], settings: StepSettings | None = None
) -> tuple[str, list[str]]:
    settings = settings or load_step("letter")
    system = [{"type": "text", "text": read_prompt("letter_system.md")}]
    content = [{"type": "text", "text": json.dumps(facts, ensure_ascii=False, indent=2)}]
    result = call_structured(
        client, "letter", case_id, settings, system, content,
        strict_schema(Justification.model_json_schema()), Justification,
    )
    warnings = []
    unknown = unsupported_numbers(result.justification, facts)
    if unknown:
        warnings.append("Uzasadnienie zawiera liczby spoza danych sprawy: " + ", ".join(unknown) + " — sprawdź.")
    return result.justification.strip(), warnings


def _or_missing(value: str, label: str) -> str:
    return value.strip() or MISSING.format(label)


def _clients_phrase(parties: LetterParties) -> str:
    if parties.clients:
        people = [f"{name} ({_or_missing(address, 'adres mocodawcy')})" for name, address in parties.clients if name]
        if people:
            return people[0] if len(people) == 1 else ", ".join(people[:-1]) + " oraz " + people[-1]
    client = _or_missing(parties.client_name, "mocodawca")
    return f"{client} ({_or_missing(parties.client_address, 'adres mocodawcy')})"


def _agreement_phrase(parties: LetterParties) -> str:
    if parties.agreement_reference.strip():
        return parties.agreement_reference.strip()
    if parties.agreement_date or parties.agreement_deed:
        date = _or_missing(parties.agreement_date, "data umowy")
        deed = f" zawartą w formie aktu notarialnego {parties.agreement_deed}" if parties.agreement_deed else ""
        return f"umową deweloperską z dnia {date}{deed}"
    return MISSING.format("data i numer umowy deweloperskiej")


def letter_values(
    parties: LetterParties, total_grosze: int, justification: str, letter_date: dt.date, config: dict[str, Any]
) -> dict[str, str]:
    firm = config.get("firm", {})
    days = parties.payment_deadline_days or config.get("payment_deadline_days", 14)
    legal = parties.legal_basis or (config.get("legal_basis") or {}).get("text") or ""
    firm_name = parties.firm_name or firm.get("name", "")
    city = parties.city or firm.get("city", "")
    lawyer = parties.lawyer_name or firm.get("lawyer", "")
    signature = ", ".join(x for x in (lawyer, parties.lawyer_title) if x)
    clients = _clients_phrase(parties)
    return {
        "draft_banner": DRAFT_BANNER,
        "firm_name": firm_name,
        "firm_address": parties.firm_address or firm.get("address", ""),
        "place_and_date": f"{city}, dnia {fmt_date(letter_date)}",
        "letter_date": fmt_date(letter_date),
        "case_number": _or_missing(parties.case_number, "sygnatura sprawy"),
        "developer_name": _or_missing(parties.developer_name, "nazwa dewelopera"),
        "developer_address": _or_missing(parties.developer_address, "adres dewelopera"),
        "client_names": _or_missing(parties.client_name, "mocodawca"),
        "clients": clients,
        "agreement": _agreement_phrase(parties),
        "agreement_date": _or_missing(parties.agreement_date, "data umowy"),
        "agreement_deed": _or_missing(parties.agreement_deed, "numer aktu notarialnego umowy"),
        "total_amount": format_pln(total_grosze),
        "payment_deadline_days": str(days),
        "bank_account": _or_missing(parties.bank_account, "numer rachunku"),
        "intro": (
            f"Działając w imieniu i na rzecz {clients}, na podstawie udzielonego "
            f"pełnomocnictwa, w związku z {_agreement_phrase(parties)}, niniejszym:"
        ),
        "demand": (
            f"wzywam do zapłaty kwoty {format_pln(total_grosze)} tytułem kar umownych za opóźnienie, "
            f"w terminie {days} dni od dnia doręczenia niniejszego wezwania, na rachunek bankowy nr "
            f"{_or_missing(parties.bank_account, 'numer rachunku')}."
        ),
        "justification": justification,
        "legal_basis": legal or MISSING.format("podstawa prawna roszczenia"),
        "closing": (
            "Brak zapłaty we wskazanym terminie spowoduje skierowanie sprawy na drogę postępowania "
            "sądowego bez dalszego wezwania."
        ),
        "lawyer": signature,
        "lawyer_name": _or_missing(lawyer, "imię i nazwisko pełnomocnika"),
        "lawyer_title": parties.lawyer_title,
        "lawyer_registry": parties.lawyer_registry,
    }


def _set_text(paragraph: Paragraph, text: str) -> None:
    runs = paragraph.runs
    if not runs:
        paragraph.add_run(text)
        return
    runs[0].text = text
    for run in runs[1:]:
        run.text = ""


def _fill_paragraph(paragraph: Paragraph, values: dict[str, str]) -> None:
    text = paragraph.text
    if "{{" not in text:
        return
    filled = _PLACEHOLDER.sub(lambda m: values.get(m.group(1), m.group(0)), text)
    chunks = [c.strip() for c in filled.split("\n\n")] if "\n\n" in filled else [filled]
    _set_text(paragraph, chunks[0])
    anchor = paragraph._p
    for chunk in chunks[1:]:
        new_p = copy.deepcopy(paragraph._p)
        anchor.addnext(new_p)
        anchor = new_p
        _set_text(Paragraph(new_p, paragraph._parent), chunk)


def render_letter(
    template: Path, out_path: Path, values: dict[str, str], rows: list[tuple[str, str, str, str]]
) -> Path:
    doc = DocxDocument(str(template))
    for table in doc.tables:
        for row in list(table.rows):
            if any("{{penalty_rows}}" in c.text for c in row.cells):
                for data in rows:
                    new_tr = copy.deepcopy(row._tr)
                    row._tr.addprevious(new_tr)
                    for cell, text in zip(_Row(new_tr, table).cells, data):
                        _set_text(cell.paragraphs[0], text)
                row._tr.getparent().remove(row._tr)
        for row in table.rows:
            for cell in row.cells:
                for p in cell.paragraphs:
                    _fill_paragraph(p, values)
    for p in list(doc.paragraphs):
        _fill_paragraph(p, values)
    out_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(out_path))
    return out_path


def penalty_rows(outcome: CalculationOutcome, capped: bool) -> list[tuple[str, str, str, str]]:
    assert outcome.result is not None
    rows = []
    for p in outcome.result.penalties:
        if p.uncapped_grosze is None:
            continue
        start = fmt_date(p.effective_deadline.deadline + dt.timedelta(days=1))
        end = fmt_date(p.end_date) + (" (trwa)" if p.ongoing else "")
        rows.append((EVENT_LABELS[p.event].capitalize(), f"{start} – {end}", str(p.days),
                     format_pln(demanded_grosze(p, capped))))
    return rows


# --- drafts (editable before rendering) and templates ---------------------------------------

# Placeholders a firm template may use. Editable ones are shown on the edit screen;
# penalty_rows is the calculator table and cannot be edited.
PLACEHOLDERS: dict[str, str] = {
    "draft_banner": "Nagłówek projektu",
    "case_number": "Sygnatura sprawy",
    "firm_name": "Kancelaria",
    "firm_address": "Adres kancelarii",
    "place_and_date": "Miejscowość i data",
    "letter_date": "Data pisma",
    "developer_name": "Deweloper",
    "developer_address": "Adres dewelopera",
    "client_names": "Mocodawcy (imiona i nazwiska)",
    "clients": "Mocodawcy z adresami",
    "agreement": "Umowa (opis do zdania)",
    "agreement_date": "Data umowy",
    "agreement_deed": "Numer aktu notarialnego umowy",
    "total_amount": "Kwota żądania",
    "payment_deadline_days": "Termin zapłaty (dni)",
    "bank_account": "Rachunek do zapłaty",
    "intro": "Wstęp (mocodawca, pełnomocnictwo, umowa)",
    "demand": "Żądanie zapłaty",
    "justification": "Uzasadnienie (projekt przygotowany przez model)",
    "legal_basis": "Podstawa prawna",
    "closing": "Zakończenie",
    "lawyer": "Podpis (prawnik, tytuł)",
    "lawyer_name": "Prawnik — imię i nazwisko",
    "lawyer_title": "Prawnik — tytuł zawodowy",
    "lawyer_registry": "Prawnik — numer wpisu",
}
TABLE_PLACEHOLDER = "penalty_rows"
REQUIRED_PLACEHOLDERS: tuple[str, ...] = ()
_MISSING_RE = re.compile(r"\[UZUPEŁNIJ: [^\]]+\]")


@dataclass
class LetterDraft:
    values: dict[str, str]
    rows: list[tuple[str, str, str, str]]
    total_grosze: int
    capped: bool
    warnings: list[str] = field(default_factory=list)

    def missing_fields(self) -> list[str]:
        return sorted({m for v in self.values.values() for m in _MISSING_RE.findall(v)})

    def to_dict(self) -> dict[str, Any]:
        return {"values": self.values, "rows": [list(r) for r in self.rows], "total_grosze": self.total_grosze,
                "capped": self.capped, "warnings": self.warnings}

    @classmethod
    def from_dict(cls, data: dict[str, Any]) -> LetterDraft:
        return cls(dict(data["values"]), [tuple(r) for r in data["rows"]], int(data["total_grosze"]),
                   bool(data["capped"]), list(data.get("warnings", [])))


def prepare_draft(
    client: LLMClient,
    case_id: str,
    ex: Extraction,
    outcome: CalculationOutcome,
    parties: LetterParties,
    *,
    capped: bool,
    letter_date: dt.date,
    config_dir: Path | None = None,
    settings: StepSettings | None = None,
    with_justification: bool = True,
) -> LetterDraft:
    """`capped` is the lawyer's choice: demand the amount with or without the contractual cap."""
    if outcome.result is None:
        raise ValueError("no calculation result: " + "; ".join(outcome.blocking))
    config = load_letter_config(config_dir)
    legal = (config.get("legal_basis") or {}).get("text") or ""
    legal = parties.legal_basis or legal
    justification, warnings = "", []
    if with_justification:
        facts = letter_facts(ex, outcome, capped, legal)
        justification, warnings = draft_justification(client, case_id, facts, settings)
    total = sum(demanded_grosze(p, capped) for p in outcome.result.penalties)
    values = letter_values(parties, total, justification, letter_date, config)
    return LetterDraft(values, penalty_rows(outcome, capped), total, capped, warnings)


def template_placeholders(template: Path) -> set[str]:
    """Placeholder names used anywhere in a .docx template (body and tables)."""
    doc = DocxDocument(str(template))
    texts = [p.text for p in doc.paragraphs]
    texts += [c.text for t in doc.tables for r in t.rows for c in r.cells]
    return {m for text in texts for m in _PLACEHOLDER.findall(text)}


def check_template(template: Path) -> tuple[list[str], list[str]]:
    """(errors, warnings) for a firm template. Raises nothing; an unreadable file is an error."""
    try:
        found = template_placeholders(template)
    except Exception:  # not a .docx / corrupted
        return ["Plik nie jest poprawnym dokumentem .docx."], []
    known = set(PLACEHOLDERS) | {TABLE_PLACEHOLDER}
    errors = [f"Brak wymaganego pola {{{{{name}}}}}." for name in REQUIRED_PLACEHOLDERS if name not in found]
    if not found & known:
        errors.append("Szablon nie zawiera żadnego pola {{...}} — system nie miałby czego uzupełnić.")
    warnings = [f"Nieznane pole {{{{{name}}}}} — zostanie w piśmie bez zmian." for name in sorted(found - known)]
    if TABLE_PLACEHOLDER not in found:
        warnings.append("Szablon nie ma tabeli z polem {{penalty_rows}} — wyliczenie nie pojawi się w tabeli.")
    if "justification" not in found:
        warnings.append("Szablon nie ma pola {{justification}} — projekt uzasadnienia nie będzie przygotowywany.")
    return errors, warnings


def render_draft(draft: LetterDraft, template: Path, out_path: Path) -> Path:
    return render_letter(template, out_path, draft.values, draft.rows)


def build_letter(
    client: LLMClient,
    case_id: str,
    ex: Extraction,
    outcome: CalculationOutcome,
    parties: LetterParties,
    out_path: Path,
    *,
    capped: bool,
    letter_date: dt.date,
    template: Path | None = None,
    config_dir: Path | None = None,
    settings: StepSettings | None = None,
) -> LetterResult:
    """Draft and render in one step (CLI). The web app drafts, lets the lawyer edit, then renders."""
    draft = prepare_draft(client, case_id, ex, outcome, parties, capped=capped, letter_date=letter_date,
                          config_dir=config_dir, settings=settings)
    path = render_draft(draft, template or TEMPLATE, out_path)
    warnings = list(draft.warnings)
    if draft.missing_fields():
        warnings.append("Pola do uzupełnienia w piśmie: " + ", ".join(draft.missing_fields()))
    return LetterResult(path, draft.total_grosze, warnings)
