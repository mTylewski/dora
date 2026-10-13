from __future__ import annotations

import datetime as dt

import pytest
from docx import Document as DocxDocument

from pipeline.calculate import calculate_from_extraction
from pipeline.clauses import ClausePattern, assess_clauses, load_library, render_library, system_prompt
from pipeline.documents import Document
from pipeline.letter import (
    LetterDraft,
    LetterParties,
    check_template,
    prepare_draft,
    render_draft,
    build_letter,
    letter_facts,
    unsupported_numbers,
)
from pipeline.settings import ROOT
from pipeline.tests.fakes import SETTINGS, FakeLLM
from pipeline.tests.test_pipeline_steps import PAGES, extraction
from pipeline.verify_quotes import Issue

PATTERNS = [
    ClausePattern("penalty-cap-percent", "Limit", "opis", "ryzyko", "", False),
    ClausePattern("penalty-exclusive", "Wyłączna", "opis", "ryzyko", "", False),
]


# --- clauses ---------------------------------------------------------------------

def test_repo_library_loads_with_unique_ids_and_unverified_sources():
    patterns = load_library()
    assert len(patterns) >= 15
    assert all(not p.verified for p in patterns if p.source in ("", "uzupełnić"))


def test_duplicate_ids_rejected(tmp_path):
    (tmp_path / "a.yaml").write_text(
        "- {id: x, title: t, pattern_description: d, why_risky: w}\n"
        "- {id: x, title: t, pattern_description: d, why_risky: w}\n", encoding="utf-8")
    with pytest.raises(ValueError):
        load_library(tmp_path)


def test_library_rendering_is_stable_and_in_prompt():
    assert render_library(PATTERNS) == render_library(list(reversed(PATTERNS)))
    prompt = system_prompt(PATTERNS)
    assert "{library}" not in prompt and "id: penalty-exclusive" in prompt


def test_assess_clauses_verifies_and_drops_unknown():
    docs = [Document("doc_1", "a.pdf", PAGES)]
    llm = FakeLLM({"flags": [
        {"clause_id": "penalty-cap-percent", "quote": "Kara 0,02% ceny", "page": 4, "document_id": "doc_1",
         "reasoning": "r", "confidence": "high"},
        {"clause_id": "penalty-exclusive", "quote": "nie ma tego w umowie", "page": 4, "document_id": None,
         "reasoning": "r", "confidence": "low"},
        {"clause_id": "made-up", "quote": "x", "page": 1, "document_id": "doc_1", "reasoning": "r",
         "confidence": "low"},
        {"clause_id": "penalty-exclusive", "quote": "UMOWA", "page": 3, "document_id": "doc_1",
         "reasoning": "r", "confidence": "low"},
        {"clause_id": "penalty-exclusive", "quote": "UMOWA", "page": 1, "document_id": "doc_9",
         "reasoning": "r", "confidence": "low"},
    ]})
    res = assess_clauses(llm, "c1", docs, PATTERNS, SETTINGS)
    assert res.dropped_unknown_ids == 1
    assert [f.needs_review for f in res.flags] == [False, True, True, True]
    assert [(i.path, i.issue) for i in res.issues] == [
        ("flags[1]", Issue.QUOTE_NOT_FOUND), ("flags[2]", Issue.PAGE_MISMATCH), ("flags[3]", Issue.UNKNOWN_DOCUMENT)]
    assert llm.calls[0]["system"][0]["cache_control"] == {"type": "ephemeral"}


# --- letter ----------------------------------------------------------------------

def outcome_and_extraction():
    ex = extraction(penalty_cap={"value": {"type": "amount", "value": 1000}, "quote": "q", "page": 1})
    return ex, calculate_from_extraction(ex, dt.date(2025, 1, 1))


def test_letter_facts_have_no_party_data_and_respect_cap_choice():
    ex, out = outcome_and_extraction()
    uncapped = letter_facts(ex, out, capped=False, legal_basis="")
    capped = letter_facts(ex, out, capped=True, legal_basis="")
    # 450 000 zł × 0,02% × 45 dni = 4 050 zł; cap 1 000 zł
    assert uncapped["total_demanded"] == "4 050,00 zł"
    assert capped["total_demanded"] == "1 000,00 zł" and capped["penalties"][0]["cap_applied"]
    assert uncapped["penalty_rate"] == "0,02% ceny za każdy dzień"
    assert uncapped["penalties"][0]["days_of_delay"] == 45


def test_letter_facts_fixed_rate_and_ongoing():
    ex = extraction(penalty_basis={"value": "fixed_per_day", "quote": "q", "page": 1},
                    penalty_rate={"value": 50, "quote": "q", "page": 1},
                    actual_handover_date={"value": None, "quote": None})
    out = calculate_from_extraction(ex, dt.date(2024, 4, 10))
    facts = letter_facts(ex, out, capped=False, legal_basis="")
    assert facts["penalty_rate"] == "50,00 zł za każdy dzień"
    assert facts["penalties"][0]["still_not_performed_as_of"] == "10.04.2024"


def test_unsupported_numbers():
    facts = {"a": "4 050,00 zł", "d": "31.03.2024", "n": 45}
    assert unsupported_numbers("Kwota 4 050,00 zł za 45 dni od 31.03.2024.", facts) == []
    assert unsupported_numbers("Kwota 5000 zł.", facts) == ["5000"]


def test_build_letter_docx(tmp_path):
    ex, out = outcome_and_extraction()
    llm = FakeLLM({"justification": "Akapit pierwszy, 45 dni.\n\nAkapit drugi, kwota 999 zł."})
    parties = LetterParties(client_name="Jan Testowy", developer_name="Deweloper Test sp. z o.o.")
    res = build_letter(llm, "c1", ex, out, parties, tmp_path / "w.docx", capped=False,
                       letter_date=dt.date(2025, 1, 2), settings=SETTINGS)
    assert res.total_grosze == 405_000
    text = "\n".join(p.text for p in DocxDocument(str(res.path)).paragraphs)
    assert "Jan Testowy" in text and "Deweloper Test sp. z o.o." in text
    assert "Akapit pierwszy" in text and "Akapit drugi" in text and "{{" not in text
    assert "PROJEKT" in text and "dnia 02.01.2025" in text
    assert "[UZUPEŁNIJ: podstawa prawna roszczenia]" in text
    table = DocxDocument(str(res.path)).tables[0]
    assert [c.text for c in table.rows[1].cells] == [
        "Opóźnienie w wydaniu lokalu", "01.04.2024 – 15.05.2024", "45", "4 050,00 zł"]
    assert len(table.rows) == 2
    assert any("999" in w for w in res.warnings)
    assert any("numer rachunku" in w for w in res.warnings)
    # party data never reaches the model
    sent = str(llm.calls[0]["messages"])
    assert "Jan Testowy" not in sent and "Deweloper Test" not in sent


def test_build_letter_requires_calculation(tmp_path):
    ex = extraction(price_gross={"value": None, "quote": None})
    out = calculate_from_extraction(ex, dt.date(2025, 1, 1))
    with pytest.raises(ValueError):
        build_letter(FakeLLM(), "c1", ex, out, LetterParties(), tmp_path / "w.docx", capped=False,
                     letter_date=dt.date(2025, 1, 1), settings=SETTINGS)


def test_template_exists():
    assert (ROOT / "templates" / "wezwanie_do_zaplaty.docx").exists()


def test_draft_roundtrip_and_render_uses_edited_values(tmp_path):
    ex, out = outcome_and_extraction()
    llm = FakeLLM({"justification": "Pierwotny akapit."})
    draft = prepare_draft(llm, "c1", ex, out, LetterParties(client_name="Jan Testowy"), capped=True,
                          letter_date=dt.date(2025, 1, 2), settings=SETTINGS)
    assert draft.total_grosze == 100_000 and draft.capped  # cap 1 000 zł
    restored = LetterDraft.from_dict(draft.to_dict())
    assert restored.rows == draft.rows and restored.values == draft.values
    assert "[UZUPEŁNIJ: numer rachunku]" in restored.missing_fields()
    restored.values["justification"] = "Akapit po edycji."
    path = render_draft(restored, ROOT / "templates" / "wezwanie_do_zaplaty.docx", tmp_path / "w.docx")
    text = "\n".join(p.text for p in DocxDocument(str(path)).paragraphs)
    assert "Akapit po edycji." in text and "Pierwotny akapit." not in text and "Jan Testowy" in text


def test_check_template_on_default_template():
    errors, warnings = check_template(ROOT / "templates" / "wezwanie_do_zaplaty.docx")
    assert errors == [] and warnings == []
