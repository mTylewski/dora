from __future__ import annotations

import datetime as dt
import re

import pyotp
import pytest
from cryptography.fernet import Fernet
from fastapi.testclient import TestClient
from sqlalchemy import select

from app.config import Settings
from app.db import AccessLog, Case, make_sessionmaker
from app.main import create_app, create_user
from app.services import Services
from pipeline.tests.fakes import FakeLLM, write_pdf
from pipeline.tests.test_anonymize import PESEL
from pipeline.tests.test_run import PAGES, extraction_payload

PASSWORD = "bardzo-dlugie-haslo-1"


def scripted_llm():
    return FakeLLM(
        {"documents": [{"document_id": "doc_1", "type": "development_agreement", "confidence": "high"}]},
        extraction_payload(),
        {"flags": [{"clause_id": "penalty-symbolic-rate", "quote": "Kara 0,02% ceny za kazdy dzien", "page": 4,
                    "document_id": "doc_1", "reasoning": "Niska stawka.", "confidence": "low"}]},
        {"justification": "Deweloper nie wydał lokalu w terminie."},
    )


@pytest.fixture
def env(tmp_path):
    settings = Settings(
        database_url=f"sqlite:///{tmp_path / 'db.sqlite'}", redis_url=None, secret_key="test-secret",
        mapping_key=Fernet.generate_key().decode(), data_dir=tmp_path / "data", retention_days=30,
        require_anon_approval=True, secure_cookies=False, max_upload_mb=5,
    )
    llm = scripted_llm()
    services = Services(settings, make_sessionmaker(settings.database_url), llm_factory=lambda: llm)
    app = create_app(services)
    client = TestClient(app)
    admin = create_user(services, "Admin@Example.com", PASSWORD, "admin")
    yield {"client": client, "services": services, "admin": admin, "llm": llm, "tmp": tmp_path}
    from app import jobs
    jobs.use_services(None)


def csrf(client, url="/login"):
    html = client.get(url).text
    return re.search(r'name="csrf" value="([^"]+)"', html).group(1)


def login(client, user, password=PASSWORD):
    token = csrf(client)
    r = client.post("/login", data={"email": user.email, "password": password, "csrf": token},
                    follow_redirects=False)
    assert r.status_code == 303 and r.headers["location"] == "/login/totp"
    token = csrf(client, "/login/totp")
    code = pyotp.TOTP(user.totp_secret).now()
    r = client.post("/login/totp", data={"code": code, "csrf": token}, follow_redirects=False)
    assert r.status_code == 303 and r.headers["location"] == "/"


def test_requires_login(env):
    r = env["client"].get("/", follow_redirects=False)
    assert r.status_code == 303 and r.headers["location"] == "/login"
    assert env["client"].get("/healthz").text == "ok"


def test_bad_password_and_throttle(env):
    c = env["client"]
    for _ in range(5):
        r = c.post("/login", data={"email": "admin@example.com", "password": "zle", "csrf": csrf(c)})
        assert "Nieprawidłowy e-mail lub hasło" in r.text
    r = c.post("/login", data={"email": "admin@example.com", "password": PASSWORD, "csrf": csrf(c)})
    assert "Zbyt wiele prób" in r.text


def test_wrong_totp(env):
    c = env["client"]
    c.post("/login", data={"email": "admin@example.com", "password": PASSWORD, "csrf": csrf(c)})
    r = c.post("/login/totp", data={"code": "000000", "csrf": csrf(c, "/login/totp")})
    assert "Nieprawidłowy kod" in r.text
    assert c.get("/", follow_redirects=False).status_code == 303


def test_csrf_required(env):
    c = env["client"]
    login(c, env["admin"])
    r = c.post("/cases/x/approve", data={"csrf": "bad"})
    assert r.status_code == 403


def test_full_case_flow(env):
    c, services, tmp = env["client"], env["services"], env["tmp"]
    login(c, env["admin"])
    pdf = write_pdf(tmp / "umowa.pdf", PAGES).read_bytes()
    token = csrf(c, "/cases/new")
    r = c.post("/cases", data={"title": "Sprawa testowa", "as_of": "2024-04-10", "csrf": token},
               files=[("files", ("umowa.pdf", pdf, "application/pdf"))], follow_redirects=False)
    assert r.status_code == 303
    case_url = r.headers["location"]
    case_id = case_url.rsplit("/", 1)[1]

    # Anonymization approval screen shows what will be replaced, nothing was sent yet
    page = c.get(case_url).text
    assert "Zatwierdzenie anonimizacji" in page and "[PESEL_1]" in page
    assert env["llm"].calls == []

    r = c.post(f"{case_url}/approve", data={"csrf": csrf(c, case_url)}, follow_redirects=False)
    assert r.status_code == 303
    sent = str([call["messages"] for call in env["llm"].calls])
    assert PESEL not in sent and "Nowakowski" not in sent

    page = c.get(case_url).text
    assert "Karta sprawy" in page and "450 000,00 zł" in page
    assert "Opóźnienie w wydaniu lokalu" in page and "900,00" in page  # 10 days × 90 zł
    assert "Symboliczna stawka kary" in page
    assert c.get(f"{case_url}/documents/doc_1").headers["content-type"] == "application/pdf"

    # Lawyer corrects the actual handover date: 01.04–15.04.2024 = 15 days × 90 zł = 1 350 zł
    r = c.post(f"{case_url}/fields", data={"csrf": csrf(c, case_url), "path": "actual_handover_date",
                                            "value": "15.04.2024"}, follow_redirects=False)
    assert r.status_code == 303
    page = c.get(case_url).text
    assert "1 350,00" in page and "wpisane przez prawnika" in page

    # Calculation date can be changed and is audited
    r = c.post(f"{case_url}/as-of", data={"csrf": csrf(c, case_url), "as_of": "2024-12-31"}, follow_redirects=False)
    assert r.status_code == 303
    assert 'value="2024-12-31"' in c.get(case_url).text

    # Invalid correction is reported, nothing saved
    r = c.post(f"{case_url}/fields", data={"csrf": csrf(c, case_url), "path": "price_gross", "value": "abc"})
    assert "nieprawidłowa liczba" in r.text

    # Confirm a field without changing it
    c.post(f"{case_url}/fields", data={"csrf": csrf(c, case_url), "path": "price_gross", "action": "confirm"})
    assert "potwierdzone" in c.get(case_url).text

    # Firm settings and a lawyer linked to the admin account
    firm_url = "/admin/firm"
    c.post(firm_url, data={"csrf": csrf(c, firm_url), "name": "Kancelaria Testowa", "address": "ul. Testowa 1",
                           "city": "Warszawa", "bank_account": "11 1111 1111 1111 1111 1111 1111",
                           "payment_deadline_days": "7", "case_number_pattern": "KAN/{rok}/{nr}", "legal_basis": ""})
    c.post(f"{firm_url}/lawyers", data={"csrf": csrf(c, firm_url), "full_name": "Anna Prawnicza",
                                        "title": "radca prawny", "registry_no": "WA-123",
                                        "user_id": str(env["admin"].id)})
    assert "Anna Prawnicza" in c.get(firm_url).text

    # Letter wizard: data suggested from documents and settings
    letter_url = f"{case_url}/letter"
    page = c.get(letter_url).text
    assert "1. Rodzaj pisma" in page and "szablon domyślny" in page
    assert 'value="Jan Nowakowski"' in page  # buyer from the komparycja, de-anonymized locally
    assert "11 1111 1111 1111 1111 1111 1111" in page and "APM Test sp. z o.o." in page
    r = c.post(f"{letter_url}/prepare", data={
        "csrf": csrf(c, letter_url), "template_id": "0", "client_include": "0",
        "client_name": ["Jan Nowakowski", ""], "client_address": ["ul. Kliencka 5, 00-001 Warszawa", ""],
        "developer_name": "APM Test sp. z o.o.", "developer_address": "ul. Budowlana 2",
        "agreement_date": "10.01.2023", "agreement_deed": "Rep. A nr 100/2023", "case_number": "KAN/2026/001",
        "lawyer_id": "1", "bank_account": "11 1111 1111 1111 1111 1111 1111", "payment_deadline_days": "7",
        "capped": "no"}, follow_redirects=False)
    assert r.status_code == 303
    letter_messages = str(env["llm"].calls[-1]["messages"])
    assert "Jan Nowakowski" not in letter_messages and "Kliencka" not in letter_messages
    page = c.get(letter_url).text
    assert "Zapisz i pobierz .docx" in page and "Jan Nowakowski (ul. Kliencka 5, 00-001 Warszawa)" in page
    assert "Deweloper nie wydał lokalu w terminie." in page
    assert "umową deweloperską z dnia 10.01.2023 zawartą w formie aktu notarialnego Rep. A nr 100/2023" in page
    assert "w terminie 7 dni" in page and "Anna Prawnicza, radca prawny" in page

    edited = "Mocodawca zawarł umowę.\n\nDeweloper nie wydał lokalu w terminie — akapit poprawiony przez prawnika."
    r = c.post(f"{letter_url}/save", data={"csrf": csrf(c, letter_url), "justification": edited,
                                           "action": "download"}, follow_redirects=True)
    assert r.headers["content-type"].startswith("application/vnd.openxmlformats")
    assert "KAN_2026_001" in r.headers["content-disposition"]
    docx_path = env["tmp"] / "pobrane.docx"
    docx_path.write_bytes(r.content)
    from docx import Document as DocxDocument
    text = "\n".join(p.text for p in DocxDocument(str(docx_path)).paragraphs)
    assert "akapit poprawiony przez prawnika" in text and "Jan Nowakowski" in text and "{{" not in text
    assert "Sygn.: KAN/2026/001" in text and "Warszawa, dnia" in text

    # A correction after preparing makes the draft stale
    c.post(f"{case_url}/fields", data={"csrf": csrf(c, case_url), "path": "actual_handover_date",
                                       "value": "20.04.2024"})
    assert "Wyliczenie w sprawie zmieniło się" in c.get(letter_url).text
    assert "1. Rodzaj pisma" in c.get(f"{letter_url}?new=1").text

    # Data is encrypted at rest
    with services.sessions() as db:
        case = db.get(Case, case_id)
        assert b"Nowakowski" not in case.prepared_enc and b"Nowakowski" not in (case.parties_enc or b"")
        assert b"Nowakowski" not in case.letter_enc
        actions = [a.action for a in db.scalars(select(AccessLog).where(AccessLog.case_id == case_id))]
    assert {"case_create", "anonymization_approved", "case_view", "correct:actual_handover_date",
            "letter_prepare:0", "letter_edit:justification", "letter_download"} <= set(actions)

    # Admin log page
    assert "letter_prepare:0" in c.get(f"/admin/log?case_id={case_id}").text

    # Delete
    r = c.post(f"{case_url}/delete", data={"csrf": csrf(c, case_url)}, follow_redirects=False)
    assert r.status_code == 303
    assert c.get(case_url).status_code == 404
    assert not services.case_dir(case_id).exists()


def test_upload_rejects_non_pdf(env):
    c = env["client"]
    login(c, env["admin"])
    r = c.post("/cases", data={"title": "x", "as_of": "2024-01-01", "csrf": csrf(c, "/cases/new")},
               files=[("files", ("a.pdf", b"hello", "application/pdf"))])
    assert "plik nie jest PDF" in r.text


def test_analysis_error_and_retry(env):
    services, tmp = env["services"], env["tmp"]
    env["llm"].responses = ["bad", "bad"]  # classification fails twice
    with services.sessions() as db:
        admin = db.get(type(env["admin"]), env["admin"].id)
    case_id = services.create_case(admin, "x", [("a.pdf", write_pdf(tmp / "a.pdf", PAGES).read_bytes())],
                                   dt.date(2024, 4, 10))
    services.approve(case_id, admin.id)
    with services.sessions() as db:
        case = db.get(Case, case_id)
        assert (case.status, case.error_code) == ("error", "invalid_output")
    env["llm"].responses = list(scripted_llm().responses)
    services.retry(case_id, admin.id)
    with services.sessions() as db:
        assert db.get(Case, case_id).status == "ready"


def test_admin_creates_user_and_lawyer_cannot_admin(env):
    c = env["client"]
    login(c, env["admin"])
    r = c.post("/admin/users", data={"csrf": csrf(c, "/admin/users"), "email": "prawnik@example.com",
                                     "password": PASSWORD, "role": "lawyer"})
    assert "otpauth://" in r.text and "<svg" in r.text
    r = c.post("/admin/users", data={"csrf": csrf(c, "/admin/users"), "email": "prawnik@example.com",
                                     "password": PASSWORD, "role": "lawyer"})
    assert "użytkownik już istnieje" in r.text

    from app.db import User
    with env["services"].sessions() as db:
        lawyer = db.scalar(select(User).where(User.email == "prawnik@example.com"))
    c.post("/logout", data={"csrf": csrf(c, "/")})
    login(c, lawyer)
    assert c.get("/admin/users").status_code == 403

    # Admin blocks the lawyer -> session no longer valid
    admin_client = TestClient(c.app)
    login(admin_client, env["admin"])
    admin_client.post(f"/admin/users/{lawyer.id}/toggle", data={"csrf": csrf(admin_client, "/admin/users")})
    assert c.get("/", follow_redirects=False).status_code == 303


def test_purge_expired(env):
    services, tmp = env["services"], env["tmp"]
    with services.sessions() as db:
        admin = db.get(type(env["admin"]), env["admin"].id)
    case_id = services.create_case(admin, "x", [("a.pdf", write_pdf(tmp / "a.pdf", PAGES).read_bytes())],
                                   dt.date(2024, 4, 10))
    assert services.purge_expired(now=dt.datetime.now() + dt.timedelta(days=1)) == 0
    assert services.purge_expired(now=dt.datetime.now() + dt.timedelta(days=31)) == 1
    with services.sessions() as db:
        assert db.get(Case, case_id) is None


def test_security_headers(env):
    r = env["client"].get("/login")
    assert "frame-ancestors 'self'" in r.headers["content-security-policy"]
    assert r.headers["x-content-type-options"] == "nosniff"


def _docx_with(paragraphs: list[str], path):
    from docx import Document as DocxDocument

    doc = DocxDocument()
    for text in paragraphs:
        doc.add_paragraph(text)
    doc.save(str(path))
    return path.read_bytes()


def test_admin_template_library(env):
    c, services, tmp = env["client"], env["services"], env["tmp"]
    login(c, env["admin"])
    url = "/admin/templates"
    assert "szablon domyślny" in c.get(url).text
    assert c.get(f"{url}/0/download").headers["content-type"].startswith("application/vnd.openxmlformats")

    r = c.post(url, data={"csrf": csrf(c, url), "name": "Zły"},
               files={"file": ("x.docx", b"not a docx", "application/octet-stream")})
    assert "nie jest poprawnym dokumentem" in r.text

    empty = _docx_with(["KANCELARIA TEST", "bez pól"], tmp / "e.docx")
    r = c.post(url, data={"csrf": csrf(c, url), "name": "Pusty"}, files={"file": ("e.docx", empty, "x")})
    assert "nie zawiera żadnego pola" in r.text

    good = _docx_with(["KANCELARIA TEST — papier firmowy", "Sygn.: {{case_number}}", "{{demand}}", "{{inne_pole}}"],
                      tmp / "g.docx")
    r = c.post(url, data={"csrf": csrf(c, url), "name": "Wezwanie kancelarii XYZ", "description": "drugie pismo"},
               files={"file": ("g.docx", good, "x")})
    assert "Dodano szablon" in r.text and "Nieznane pole {{inne_pole}}" in r.text
    assert "projekt uzasadnienia nie będzie przygotowywany" in r.text
    assert "Wezwanie kancelarii XYZ" in c.get(url).text
    assert c.get(f"{url}/1/download").status_code == 200

    c.post(f"{url}/1/deactivate", data={"csrf": csrf(c, url)})
    assert "Wezwanie kancelarii XYZ" not in c.get(url).text
    assert c.get(f"{url}/1/download").status_code == 404


def test_case_numbers_follow_the_firm_pattern(env):
    services, tmp = env["services"], env["tmp"]
    from app import firm
    with services.sessions() as db:
        admin = db.get(type(env["admin"]), env["admin"].id)
        firm.get_settings(db).case_number_pattern = "SP/{rok}/{nr}"
        db.commit()
    pdf = write_pdf(tmp / "a.pdf", PAGES).read_bytes()
    first = services.create_case(admin, "a", [("a.pdf", pdf)], dt.date(2024, 4, 10))
    second = services.create_case(admin, "b", [("a.pdf", pdf)], dt.date(2024, 4, 10))
    manual = services.create_case(admin, "c", [("a.pdf", pdf)], dt.date(2024, 4, 10), case_number="WŁASNA/7")
    year = dt.date.today().year
    with services.sessions() as db:
        numbers = [db.get(Case, cid).case_number for cid in (first, second, manual)]
    assert numbers == [f"SP/{year}/001", f"SP/{year}/002", "WŁASNA/7"]


def test_saving_unchanged_value_counts_as_confirmation(env):
    services, tmp = env["services"], env["tmp"]
    with services.sessions() as db:
        admin = db.get(type(env["admin"]), env["admin"].id)
    case_id = services.create_case(admin, "x", [("a.pdf", write_pdf(tmp / "a.pdf", PAGES).read_bytes())],
                                   dt.date(2024, 4, 10))
    services.approve(case_id, admin.id)
    services.correct(case_id, "price_gross", "450000,00", admin.id)   # same as the document
    services.correct(case_id, "penalty_rate", "0,0003", admin.id)      # changed
    with services.sessions() as db:
        view = services.case_view(db.get(Case, case_id))
    fields = {f["path"]: f for f in view["fields"]}
    assert fields["price_gross"]["corrected"] and not fields["price_gross"]["manual"]
    assert fields["penalty_rate"]["manual"]

    services.revert(case_id, "penalty_rate", admin.id)
    with services.sessions() as db:
        view = services.case_view(db.get(Case, case_id))
    rate = {f["path"]: f for f in view["fields"]}["penalty_rate"]
    assert not rate["corrected"] and rate["value"] == "0.0002"


def test_progress_is_shown_while_processing(env):
    c, services, tmp = env["client"], env["services"], env["tmp"]

    class Paused:  # background jobs not started yet
        def enqueue(self, job, case_id):
            pass

    services.queue = Paused()
    login(c, env["admin"])
    with services.sessions() as db:
        admin = db.get(type(env["admin"]), env["admin"].id)
    case_id = services.create_case(admin, "x", [("a.pdf", write_pdf(tmp / "a.pdf", PAGES).read_bytes())],
                                   dt.date(2024, 4, 10))
    services._progress(case_id, "Dokument 1 z 1 — OCR: strona 3 z 27", 40)
    page = c.get(f"/cases/{case_id}").text
    assert "spinner" in page and '<progress class="progress" max="100" value="40">' in page
    assert "OCR: strona 3 z 27" in page and 'http-equiv="refresh" content="3"' in page
