"""Create the default demand letter template (templates/wezwanie_do_zaplaty.docx).

The law firm can replace it with its own letterhead; keep the {{placeholders}}.
Each placeholder must sit in its own paragraph or table cell text.

Usage: python scripts/make_letter_template.py
"""

from __future__ import annotations

import sys
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))

TARGET = ROOT / "templates" / "wezwanie_do_zaplaty.docx"


def build(path: Path = TARGET) -> Path:
    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "Times New Roman"
    style.font.size = Pt(12)

    banner = doc.add_paragraph("{{draft_banner}}")
    banner.runs[0].bold = True
    banner.alignment = WD_ALIGN_PARAGRAPH.CENTER

    doc.add_paragraph("{{firm_name}}")
    doc.add_paragraph("{{firm_address}}")
    doc.add_paragraph("Sygn.: {{case_number}}")
    right = doc.add_paragraph("{{place_and_date}}")
    right.alignment = WD_ALIGN_PARAGRAPH.RIGHT

    doc.add_paragraph("")
    doc.add_paragraph("{{developer_name}}")
    doc.add_paragraph("{{developer_address}}")
    doc.add_paragraph("")

    title = doc.add_paragraph("PRZEDSĄDOWE WEZWANIE DO ZAPŁATY")
    title.runs[0].bold = True
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER

    doc.add_paragraph("{{intro}}")
    doc.add_paragraph("{{demand}}")
    doc.add_paragraph("Wyliczenie kary umownej:")

    table = doc.add_table(rows=1, cols=4)
    table.style = "Table Grid"
    for cell, text in zip(table.rows[0].cells, ("Tytuł", "Okres opóźnienia", "Dni", "Kwota")):
        cell.text = text
    row = table.add_row().cells
    row[0].text = "{{penalty_rows}}"

    heading = doc.add_paragraph("Uzasadnienie")
    heading.runs[0].bold = True
    doc.add_paragraph("{{justification}}")
    doc.add_paragraph("{{legal_basis}}")
    doc.add_paragraph("{{closing}}")
    doc.add_paragraph("")
    sign = doc.add_paragraph("{{lawyer}}")
    sign.alignment = WD_ALIGN_PARAGRAPH.RIGHT

    path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(path)
    return path


if __name__ == "__main__":
    print(f"written {build().relative_to(ROOT)}")
